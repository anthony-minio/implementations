#!/usr/bin/env python3
"""Restyle pandoc's .docx output to match the Ramp Configuration Reference.

Tables with a header row get a navy header with white text. Tables whose header
row is empty (written as `| | |` in markdown) come out of pandoc with no header
row; those become key/value tables with the first column shaded grey.
"""
import sys

import docx
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

NAVY = "1F3B5C"
LABEL_GREY = "F2F4F7"
BORDER = "BFC7D1"
BODY_FONT = "Arial"
CONTENT_TWIPS = 9360  # 6.5in of text width on Letter with 1in margins


def set_cell_fill(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in tc_pr.findall(qn("w:shd")):
        tc_pr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def set_borders(table):
    tbl_pr = table._tbl.tblPr
    for old in tbl_pr.findall(qn("w:tblBorders")):
        tbl_pr.remove(old)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), BORDER)
        borders.append(el)
    tbl_pr.append(borders)


def set_widths(table, fractions):
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(CONTENT_TWIPS))
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)
    grid = table._tbl.tblGrid
    widths = [int(CONTENT_TWIPS * f) for f in fractions]
    for col, w in zip(grid.findall(qn("w:gridCol")), widths):
        col.set(qn("w:w"), str(w))
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            tc_w = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            tc_w.type = "dxa"
            tc_w.w = w


def column_fractions(table, kv):
    if kv and len(table.columns) == 2:
        return [0.28, 0.72]
    # Share width by the longest text in each column, with a floor so short columns stay readable.
    longest = [max(len(c.text) for c in col.cells) or 1 for col in table.columns]
    longest_word = [max(len(w) for c in col.cells for w in c.text.split() or [""]) for col in table.columns]
    weights = [max(min(n, 90), 3.0 * w, 16) for n, w in zip(longest, longest_word)]
    total = sum(weights)
    fractions = [w / total for w in weights]
    # Keep a single bold word, such as a header over an empty column, on one line.
    floor = 0.10
    short = [f < floor for f in fractions]
    spare = 1 - floor * sum(short)
    rest = sum(f for f, s in zip(fractions, short) if not s)
    return [floor if s else f * spare / rest for f, s in zip(fractions, short)]


def style_runs(cell, size=Pt(10), bold=None, color=None):
    for p in cell.paragraphs:
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        for r in p.runs:
            r.font.size = size
            r.font.name = BODY_FONT
            if bold is not None:
                r.font.bold = bold
            if color is not None:
                r.font.color.rgb = RGBColor.from_string(color)


def style_table(table):
    # Pandoc drops an all-empty header row, so a table without a marked header is key/value.
    first = table.rows[0]._tr
    kv = first.trPr is None or first.trPr.find(qn("w:tblHeader")) is None
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_borders(table)
    set_widths(table, column_fractions(table, kv))
    for i, row in enumerate(table.rows):
        for j, cell in enumerate(row.cells):
            if not kv and i == 0:
                set_cell_fill(cell, NAVY)
                style_runs(cell, bold=True, color="FFFFFF")
            elif kv and j == 0:
                set_cell_fill(cell, LABEL_GREY)
                style_runs(cell, bold=True)
            else:
                style_runs(cell)


def style_fonts(doc):
    # Look styles up by display name; doc.styles[...] mistranslates pandoc's "Heading 1".
    by_name = {s.name: s for s in doc.styles}
    for name in ("Normal", "Body Text", "First Paragraph", "Compact"):
        if name in by_name:
            st = by_name[name]
            st.font.name = BODY_FONT
            st.font.size = Pt(11)
            st.font.color.rgb = RGBColor(0, 0, 0)
    spec = {
        "Title": (Pt(28), "000000", BODY_FONT),
        "Subtitle": (Pt(16), "666666", "Georgia"),
        "Heading 1": (Pt(16), "2E74B5", BODY_FONT),
        "Heading 2": (Pt(13), "2E74B5", BODY_FONT),
        "Heading 3": (Pt(12), "1F4D78", BODY_FONT),
        "Image Caption": (Pt(9), "555555", BODY_FONT),
    }
    for name, (size, color, font) in spec.items():
        if name in by_name:
            st = by_name[name]
            st.font.size = size
            st.font.name = font
            st.font.color.rgb = RGBColor.from_string(color)
            st.font.italic = name in ("Subtitle", "Image Caption")
            if name.startswith("Heading"):
                st.font.bold = True
    for name in ("Title", "Subtitle"):
        if name in by_name:
            by_name[name].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT


def style_images(doc):
    for p in doc.paragraphs:
        if p.style.name in ("Captioned Figure", "Figure") or p._p.findall(".//" + qn("w:drawing")):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if p.style.name == "Image Caption":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER


def main(path):
    doc = docx.Document(path)
    for section in doc.sections:
        section.page_width, section.page_height = Inches(8.5), Inches(11)
        section.left_margin = section.right_margin = Inches(1)
        section.top_margin = section.bottom_margin = Inches(1)
    style_fonts(doc)
    style_images(doc)
    for table in doc.tables:
        style_table(table)
    doc.save(path)


if __name__ == "__main__":
    main(sys.argv[1])
