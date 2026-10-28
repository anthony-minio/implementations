#!/usr/bin/env python3
# Salesforce matches by email domain, for review: loaded customers with no Salesforce match whose email domain is a
# Salesforce account's website. Each row carries a recommendation; nothing is applied to the load until Anthony decides.
#   python3 tools/build_sf_review.py   (run from the customers folder after build_review.py)
import collections, csv, difflib, re
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

OUT = "data/MinIO_Customer_Salesforce_Matches_to_Review_20261028.xlsx"
rows = list(csv.DictReader(open("data/customer-load-draft.csv", encoding="utf-8-sig")))
load = [r for r in rows if not r["Duplicate Action"].startswith("Merge into")]
srows = list(openpyxl.load_workbook("data/sfdc-accounts-customers-20261025.xlsx", read_only=True).active.iter_rows(values_only=True))
hdr = next(r for r in srows if r and "Account ID" in r)
accounts = [{k: "" if v in (None, "-") else str(v) for k, v in zip(hdr, r)} for r in srows
            if r and isinstance(r[2], str) and re.match(r"001\w{12}$", r[2])]

def id18(i):
    alpha = "ABCDEFGHIJKLMNOPQRSTUVWXYZ012345"
    return i + "".join(alpha[sum(1 << n for n, ch in enumerate(i[k:k + 5]) if ch.isupper())] for k in (0, 5, 10))

used = {r["Salesforce Account ID"][:15]: r["Customer ID"] for r in load if r["Salesforce Account ID"]}
by_site = collections.defaultdict(list)
for a in accounts:
    site = re.sub(r"^www\.", "", re.sub(r"^https?://", "", a["Website"].lower())).split("/")[0]
    if site:
        by_site[site].append(a)

# Recommendations from reading each candidate against the Sage record (2026-10-28).
REJECT = {("C-00030", "AT&T"): "AT&T US parent, no Subnet ID; the customer is AT&T Mexico",
          ("C-00100", "IBM India"): "IBM India, not IBM Corporation (US)",
          ("C-00170", "Herzum Inc [C]"): "Herzum is the reseller's domain, not the customer's",
          ("C-00214", "NASA - AMES"): "Another NASA center; JPL is not in Salesforce under this domain",
          ("C-00214", "NASA GSFC"): "Another NASA center; JPL is not in Salesforce under this domain",
          ("C-00277", "RundfunkBerlin-Brandenburg"): "Different organisation sharing the ivz.de domain",
          ("C-00395", "USACE"): "army.mil is shared by the whole US Army; not Yuma Proving Ground",
          ("C-00395", "US Army"): "army.mil is shared by the whole US Army; not Yuma Proving Ground",
          ("C-00286", "US Army"): "USACE has its own account (Subnet 629)",
          ("C-00172", "Orange-RDC"): "Orange RDC is C-00276; Orange-GOS is this customer",
          ("C-00512", "Xerox"): None}  # None: decided by Subnet ID below
GROUP = {"C-00050", "C-00090", "C-00127", "C-00159", "C-00211", "C-00236", "C-00244", "C-00256", "C-00279", "C-00289",
         "C-00310", "C-00540", "C-00558", "C-00590", "C-00612"}
CONFLICT = {
    "C-00219": "Salesforce account already matched to C-00120 Federal Home Loan Mortgage Corporation: Freddie Mac is the same company, so C-00219 and C-00120 may be duplicates",
    "C-00303": "Equinix (Subnet 691) also matches C-00364 Equinix (US) Enterprises: one Subnet account, two Sage customers",
    "C-00364": "Equinix (Subnet 691) also matches C-00303 Equinix Services: one Subnet account, two Sage customers",
    "C-00167": "Disney Entertainment & Sports (Subnet 385) also matches C-00218: one Subnet account, two Sage customers",
    "C-00218": "Disney Entertainment & Sports (Subnet 385) also matches C-00167: one Subnet account, two Sage customers",
    "C-00208": "Viennalife (Subnet 329) also fits C-00097 Aegon Emeklilik (same address): C-00097 may be a duplicate of C-00208",
    "C-00116": "ZATCA has four Subnet accounts in Salesforce (371, 715, 731, 852; 852 is matched to C-00445): which is this customer?",
    "C-00516": "Leidos has three accounts in Salesforce (Leidos-1 Subnet 895, Leidos-2 Subnet 972, Leidos-3 no Subnet): which is this customer?",
    "C-00640": "Hollister has two Subnet accounts in Salesforce (1012 Hollister - Credit Card, 1132 Hollister Inc.): which is this customer?",
}
norm = lambda s: re.sub(r"[^a-z0-9 ]", " ", s.lower())
out = []
for r in load:
    if r["Salesforce Account ID"]:
        continue
    domains = {e.split("@")[1].lower() for e in (r["Customer Invoice Email"] + "; " + r["Additional Emails"]).split("; ")
               if "@" in e} - {"gmail.com", "yahoo.com", "hotmail.com", "outlook.com"}
    seen = set()
    for d in sorted(domains):
        for site, accs in by_site.items():
            if site != d and not d.endswith("." + site):
                continue
            for a in accs:
                if a["Account ID"] in seen:
                    continue
                seen.add(a["Account ID"])
                key = (r["Customer ID"], a["Account Name"])
                if key in REJECT and REJECT[key]:
                    rec, why = "Reject", REJECT[key]
                elif r["Customer ID"] == "C-00512":
                    rec, why = ("Same company: accept", "Xerox France account with a Subnet ID; the contact is French") if a["Subnet Account Id"] \
                        else ("Reject", "Xerox US head office, no Subnet ID")
                elif r["Customer ID"] in CONFLICT:
                    rec, why = "Decide", CONFLICT[r["Customer ID"]]
                elif a["Account ID"] in used:
                    rec, why = "Decide", f"Salesforce account already matched to {used[a['Account ID']]}"
                elif r["Customer ID"] in GROUP:
                    rec, why = "Same group: decide", "Parent or affiliate of the customer, not the same legal entity"
                else:
                    rec, why = "Same company: accept", ""
                out.append([r["Customer ID"], r["Company Name"], r["City"], r["Country"], d, a["Account Name"], id18(a["Account ID"]),
                            a["Subnet Account Id"], a["Billing City"], a["Billing Country"], used.get(a["Account ID"], ""), rec, why, ""])
# Aegon Emeklilik has no email in Sage, so no domain; it was found by name (Viennalife, formerly Aegon).
via = next(a for a in accounts if a["Account Name"].startswith("Viennalife"))
c97 = next(r for r in load if r["Customer ID"] == "C-00097")
out.append(["C-00097", c97["Company Name"], c97["City"], c97["Country"], "(name: formerly Aegon)", via["Account Name"],
            id18(via["Account ID"]), via["Subnet Account Id"], via["Billing City"], via["Billing Country"], "", "Decide",
            "Same account fits C-00208 Aegon_Viennalife (same address): C-00097 may be a duplicate of C-00208", ""])

thin = Side(style="thin", color="D6DEE3")
grid = Border(left=thin, right=thin, top=thin, bottom=thin)
FILL = {"Same company: accept": "DCEFEF", "Same group: decide": "FFF2CF", "Decide": "FFF2CF", "Reject": "FCE5E2"}
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Matches to review"
ws["A1"] = ("Salesforce accounts found by email domain for customers in the load with no Salesforce match. Accepting a row sets "
            "the customer's Salesforce Account ID and Subnet ID. Fill in Your decision (Accept or Reject); nothing is applied until then.")
ws.merge_cells("A1:N1"); ws["A1"].alignment = Alignment(wrap_text=True); ws.row_dimensions[1].height = 32
heads = ["Customer ID", "Customer (Sage)", "City", "Country", "Matched on", "Salesforce account", "Salesforce Account ID",
         "Subnet ID", "Salesforce city", "Salesforce country", "Already matched to", "Recommendation", "Why", "Your decision"]
for k, h in enumerate(heads, 1):
    c = ws.cell(row=2, column=k, value=h); c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="E8EEF1"); c.border = grid
order = {"Decide": 0, "Same group: decide": 1, "Reject": 2, "Same company: accept": 3}
for i, row in enumerate(sorted(out, key=lambda x: (order[x[11]], x[0])), 3):
    for k, v in enumerate(row, 1):
        c = ws.cell(row=i, column=k, value=v or None); c.border = grid; c.alignment = Alignment(vertical="top", wrap_text=k in (2, 6, 13))
    ws.cell(row=i, column=12).fill = PatternFill("solid", fgColor=FILL[row[11]])
for k, w in enumerate([10, 34, 16, 14, 22, 34, 20, 9, 16, 14, 12, 20, 60, 14], 1):
    ws.column_dimensions[openpyxl.utils.get_column_letter(k)].width = w
ws.freeze_panes = "C3"
ws.auto_filter.ref = f"A2:N{len(out) + 2}"
wb.save(OUT)
print(len(out), "candidate rows for", len({x[0] for x in out}), "customers:", collections.Counter(x[11] for x in out), "->", OUT)

# The interactive review page (data/, not committed): one question per row, what Accept changes, and a recommendation.
# Decisions save to the page's database. Recommendations reviewed row by row on 2026-10-28.
import json
info = {}
h = open("data/customer-review.html", encoding="utf-8").read()
k0 = h.index("DATA =") + 6
D, _ = json.JSONDecoder().raw_decode(h[k0:].lstrip())
cix = {c: k for k, c in enumerate(D["cols"])}
for x in D["rows"]:
    rv = dict(zip(D["review"], x["review"]))
    info[x["cells"][cix["Customer ID"]][0]] = {"last": rv["Last invoice"], "balance": rv["Balance"],
                                               "name": x["cells"][cix["Company Name"]][0]}
# Specific calls (customer, Salesforce account name prefix) -> (recommendation, confidence, reason).
CALL = {
    ("C-00208", "Viennalife"): ("accept", "high", "Same company: Aegon's Turkish pension business, renamed Viennalife; same address and email domain. C-00208 is the record invoiced in 2026."),
    ("C-00219", "Federal Home Loan"): ("accept", "high", "Freddie Mac is the Federal Home Loan Mortgage Corporation. C-00219 is the record invoiced (2026-09-28, 120,000.00 open), so it should carry the Subnet ID; see the merge question for C-00120."),
    ("C-00364", "Equinix"): ("accept", "high", "Equinix as a customer (Sage type Direct Sales, 279,552.00 open). Subnet 691 belongs here."),
    ("C-00303", "Equinix"): ("reject", "high", "In Sage this record is Equinix as a billing vendor (category Other), a different relationship from Equinix the customer. Subnet 691 goes to C-00364."),
    ("C-00167", "Disney Entertainment"): ("accept", "high", "Sage bills this customer to Disney Entertainment & Sports LLC, the Salesforce account's own name."),
    ("C-00218", "Disney Entertainment"): ("reject", "high", "Disney Parks Technology Services is a different Disney company; Subnet 385 goes to C-00167, which Sage bills to Disney Entertainment & Sports."),
    ("C-00640", "Hollister Inc"): ("accept", "high", "The invoiced Hollister customer (billed through SHI, 43,200.00 open) matches Hollister Inc., Subnet 1132."),
    ("C-00116", "ZATCA - Masterworks"): ("accept", "medium", "ZATCA has four Subnet accounts. Paired by date as for L3Harris: the earlier Sage customer with the earliest Subnet account (371); C-00445 already has 852. Confirm with Nisha."),
    ("C-00116", "ZATCA-Certscan"): ("reject", "medium", "Not this customer's Subnet account if 371 is. Subnet 715 has no Sage customer of its own: may need a new customer, as for L3Harris. Ask Nisha."),
    ("C-00116", "ZATCA-Rapiscan"): ("reject", "medium", "Not this customer's Subnet account if 371 is. Subnet 731 has no Sage customer of its own: may need a new customer, as for L3Harris. Ask Nisha."),
    ("C-00116", "ZATCA-CCC"): ("reject", "high", "Already matched to C-00445 ZATCA-Monitoring."),
    ("C-00516", "Leidos-1"): ("accept", "medium", "Leidos has three Salesforce accounts. Paired by date: the first (Subnet 895). Confirm with Nisha."),
    ("C-00516", "Leidos-2"): ("reject", "medium", "Not this customer's Subnet account if 895 is. Subnet 972 has no Sage customer of its own: may need a new customer. Ask Nisha."),
    ("C-00516", "Leidos-3"): ("reject", "medium", "No Subnet ID; not needed if Leidos-1 is accepted. Ask Nisha."),
    ("C-00512", "Xerox"): None,  # decided by Subnet ID below
}
GROUP_NOTE = "Salesforce holds this under the group or parent company's name; the Sage customer is the billed entity in that group, and the Subnet account is its deal."
def call(row):
    cid, sfname, subnet, rec0, why0 = row[0], row[5], row[7], row[11], row[12]
    for (c, pre), v in CALL.items():
        if c == cid and sfname.startswith(pre) and v:
            return v
    if cid == "C-00512":
        return ("accept", "medium", "Xerox France account with a Subnet ID; the Sage contact is French. Sage's name is a Xerox service, not a legal entity.") if subnet \
            else ("reject", "high", "Xerox US head office with no Subnet ID; the France account is the match.")
    if rec0 == "Reject":
        return ("reject", "high", why0)
    if rec0 == "Same group: decide":
        return ("accept", "medium", GROUP_NOTE) if subnet else ("reject", "high", "Parent or group company with no Subnet ID: it would link a different legal entity and add nothing.")
    same_ctry = row[3] and row[9] and (row[3] == row[9] or (row[3], row[9]) in {("Türkiye", "Turkey"), ("Côte d'Ivoire", "Cote d'Ivoire"), ("Korea (the Republic of)", "Korea, Republic of"), ("Czechia", "Czech Republic")})
    return ("accept", "high", "Same company: the names match and the Salesforce website is the customer's email domain"
            + (", in the same country." if same_ctry else f"; Salesforce lists the head office in {row[9] or 'another country'}."))
items = []
for row in out:
    cid = row[0]
    if cid == "C-00097":
        continue  # asked as a merge question below
    rec, conf, why = call(row)
    sub = row[7] or "none"
    effect = (f"{cid} gets Salesforce Account ID {row[6]} and Subnet ID {row[7]}." if row[7] else f"{cid} gets Salesforce Account ID {row[6]} (the account has no Subnet ID).")
    items.append({"kind": "match", "cid": cid, "cname": row[1], "cplace": ", ".join(x for x in row[2:4] if x),
                  "last": info.get(cid, {}).get("last", ""), "balance": info.get(cid, {}).get("balance", ""),
                  "q": f"Is the Salesforce account \u201c{row[5]}\u201d the same company as this customer?",
                  "sf": f"{row[5]} \u00b7 Subnet {sub} \u00b7 {', '.join(x for x in row[8:10] if x) or 'no address'} \u00b7 website {row[4]}",
                  "sfid": row[6], "subnet": row[7], "effect": effect, "rec": rec, "conf": conf, "why": why})
def extra(cid, kind, q, sf, effect, rec, conf, why, key):
    items.append({"kind": kind, "cid": cid, "cname": info[cid]["name"], "cplace": "", "last": info[cid]["last"],
                  "balance": info[cid]["balance"], "q": q, "sf": sf, "sfid": "", "subnet": "", "effect": effect,
                  "rec": rec, "conf": conf, "why": why, "key": key})
extra("C-00097", "merge", "Is C-00097 Aegon Emeklilik the same company as C-00208 Aegon_Viennalife?",
      "C-00208 Aegon_Viennalife \u00b7 last invoiced 2026-08-17 \u00b7 Salesforce Viennalife (formerly Aegon), Subnet 329",
      "C-00097 is merged into C-00208 and not loaded (617 becomes 616).", "accept", "high",
      "Sage's own name for C-00208, Aegon_Viennalife, links the two: Aegon Emeklilik is the company's name before it became Viennalife. "
      "C-00097 has nothing else in Sage: no address, contact, email or invoice. (The address it shows in the load was copied from "
      "Salesforce's Viennalife account on 2026-10-28, so it is not separate evidence.)",
      "C-00097__merge__C-00208")
extra("C-00120", "merge", "Is C-00120 Federal Home Loan Mortgage Corporation the same company as C-00219 Freddie Mac?",
      "C-00219 Freddie Mac - Federal Home Loan Mortgage Corporation \u00b7 last invoiced 2026-09-28, 120,000.00 open",
      "C-00120 is merged into C-00219 and not loaded; C-00219 takes Subnet ID 211 (617 becomes 616).", "accept", "high",
      "Freddie Mac is the Federal Home Loan Mortgage Corporation. C-00120 was never invoiced; C-00219 is the active record.",
      "C-00120__merge__C-00219")
extra("C-00640", "unmerge", "Is the Salesforce account \u201cHollister - Credit Card\u201d (Subnet 1012) the Chargebee record C-00575 Hollister.com?",
      "C-00575 Hollister.com \u00b7 Chargebee, never invoiced in Sage \u00b7 now merged into C-00640",
      "C-00575 is loaded as its own customer with Subnet ID 1012 instead of being merged into C-00640 (617 becomes 618).",
      "accept", "medium",
      "A credit-card account is how Chargebee portal customers pay. Each Subnet account is its own customer, so the Chargebee record should not stay merged into the invoiced Hollister customer.",
      "C-00575__unmerge__1012")
for it in items:
    it.setdefault("key", (it["cid"] + "__" + (it["sfid"] or re.sub(r"[^A-Za-z0-9]", "", it["sf"][:40])))[:190])
items = [it for it in items if not (it["cid"] == "C-00640" and it["kind"] == "match" and "Credit Card" in it["sf"])]
rank = {"medium": 0, "high": 1}
items.sort(key=lambda it: (rank[it["conf"]] if it["kind"] == "match" and it["rec"] == "accept" else 0 if it["kind"] != "match" else 1,
                           it["rec"] != "reject", it["cid"]))
tpl = open("tools/sf-review-template.html", encoding="utf-8").read()
open("data/sf-match-review.html", "w", encoding="utf-8").write(
    tpl.replace("/*DATA*/null", json.dumps(items, ensure_ascii=False).replace("</", "<\\/")))
print("review page ->", len(items), "questions;", collections.Counter((it["kind"], it["rec"], it["conf"]) for it in items))
