#!/usr/bin/env python3
# Build the vendor load workbook for SOA, the controller and AP: the notes on one tab, the load table on another.
# Full version: tax IDs and bank details in full (Anthony, 2026-10-28). Never commit the output.
#   python3 tools/build_delivery.py   (run from the vendors folder after build_review.py --full)
import collections, json
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

OUT = "data/MinIO_Vendor_Load_to_NetSuite_20261028.xlsx"
ASANA = "https://app.asana.com/1/163565154279805/project/1219443754801875/task/1219707334098551"  # FS-51
s = open("data/vendor-review-full.html", encoding="utf-8").read()
data = json.loads(s[s.index("const DATA = ") + 13:s.index(";\nconst $")])
cols, review, required = data["cols"], data["review"], set(data["required"])
rows, B = data["rows"], data["bridge"]

# Same colors as the review page (light theme).
FILL = {"ramp": "E3EDF9", "rule": "DCEFEF", "prop": "FFF2CF", "miss": "FCE5E2"}
INK = {"ramp": "2B64A8", "rule": "0D6B70", "prop": "8A5D00", "miss": "A8261F"}
GROUPS = [("Identity", 9), ("Contact and delivery", 8), ("Address 1", 12), ("Address 2", 12), ("Financial", 18),
          ("Source", 2), ("Added columns", 7), ("Bank details", 5), ("Review (not loaded)", len(review) + 1)]
assert sum(n for _, n in GROUPS) == len(cols) + len(review) + 1
thin = Side(style="thin", color="D6DEE3")
grid = Border(left=thin, right=thin, top=thin, bottom=thin)
head_fill = PatternFill("solid", fgColor="E8EEF1")
wrap_top = Alignment(wrap_text=True, vertical="top")

wb = openpyxl.Workbook()
ns = wb.active
ns.title = "Notes"
ws = wb.create_sheet("Vendors-Draft")
col = 1
for name, span in GROUPS:
    ws.cell(row=1, column=col, value=name.upper())
    ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + span - 1)
    for k in range(col, col + span):
        c = ws.cell(row=1, column=k); c.fill = head_fill; c.border = grid
        c.font = Font(bold=True, size=9, color="5B6874")
    col += span
headers = cols + review + ["Flags"]
for k, h in enumerate(headers, 1):
    c = ws.cell(row=2, column=k, value=h + (" *" if h in required else ""))
    c.font = Font(bold=True); c.fill = head_fill; c.border = grid
    c.alignment = Alignment(vertical="top", wrap_text=True)
for i, r in enumerate(rows, 3):
    vals = [(v, src) for v, src in r["cells"]] + [(v, "") for v in r["review"]] + [("; ".join(r["flags"]), "")]
    for k, (v, src) in enumerate(vals, 1):
        c = ws.cell(row=i, column=k, value=v if v != "" else None)
        c.number_format = "@"
        c.border = grid
        if src in FILL:
            c.fill = PatternFill("solid", fgColor=FILL[src])
            c.font = Font(color=INK[src], italic=(src == "prop"))
        if k > len(cols):
            c.fill = PatternFill("solid", fgColor="F3F6F7") if not c.fill.fgColor.rgb or c.fill.fgColor.rgb == "00000000" else c.fill
widths = {"Company Name": 40, "Addressee": 32, "Address 1": 36, "Address 2": 24, "Email": 32, "Legal Name": 32,
          "Comments": 30, "Flags": 70, "Question Categories": 26, "Duplicate Action": 34, "Merged Sage IDs": 18,
          "Bank Name": 28, "Default Expense Account": 30, "Default Payables Account": 30, "Ramp Vendor ID": 38}
for k, h in enumerate(headers, 1):
    ws.column_dimensions[get_column_letter(k)].width = widths.get(h, 16)
ws.row_dimensions[2].height = 30
ws.freeze_panes = "H3"
ws.auto_filter.ref = f"A2:{get_column_letter(len(headers))}{len(rows) + 2}"

# Figures for the notes, counted from the table itself.
ix = {c: k for k, c in enumerate(cols)}
V = lambda r, c: r["cells"][ix[c]][0]
QC = lambda r: V(r, "Question Categories").split("; ")
has_flag = lambda r, p: any(f.startswith(p) for f in r["flags"])
merged = [r for r in rows if V(r, "Duplicate Action").startswith("Merge into")]
archived = [r for r in rows if V(r, "Duplicate Action").startswith("Archive in Ramp")]
gone = merged + archived
nsv = [r for r in rows if r not in gone]  # vendors in NetSuite
by_sync = [r for r in nsv if "Sage only" not in QC(r)]
from_file = [r for r in nsv if "Sage only" in QC(r)]
sage_only = [r for r in rows if "Sage only" in QC(r)]
kept_dups = [r for r in nsv if V(r, "Merged Sage IDs") and not V(r, "Duplicate Action").startswith("Keep, merges")]
bank = collections.Counter(V(r, "Bank Source") for r in nsv if V(r, "Bank Account Number"))
bank_differ = [r for r in nsv if has_flag(r, "Bank account differs")]
split_names = [r for r in nsv if has_flag(r, "First Name and Last Name taken")]
no_names = [r for r in nsv if has_flag(r, "Individual without first and last name")]
no_coa = [r for r in nsv if has_flag(r, "Sage expense account")]
ids = lambda rs: ", ".join(f"{V(r, 'Legacy Vendor ID') or V(r, 'Ramp Vendor ID')} {V(r, 'Company Name')}" for r in rs)
synced_id = [r for r in nsv if V(r, "Vendor ID") == "SYNCED FROM RAMP"]
assert len(synced_id) == len(by_sync)
n = lambda x: f"{x:,}"
assert len(nsv) == len(by_sync) + len(from_file) and len(rows) == B["active"] - B["out_of_scope"] - B["employees_out"] + len(B["ramp_only"])

DECISIONS = [
    ("Sage only vendors",
     f"{n(B['sage_only'])} active Sage vendors are not in Ramp. {B['recent']} have activity since 2024-10-26, "
     f"{B['older']} have older activity only, and {B['sage_only_none']} have no activity recorded.",
     "Load only the vendors with activity since 2024-10-26 (Ming, meeting of 2026-10-28).",
     f"The other {n(B['out_of_scope'])} are not in the file. Of the {B['recent']}, {B['employees_out']} MinIO employees "
     f"are also left out (next decision), so {B['sage_only_recent']} Sage only vendors are on the Vendors-Draft tab, marked "
     "\"Sage only\" in Question Categories."),
    ("MinIO employees",
     f"{B['employees_out']} of the Sage only vendors with recent activity are MinIO employees, set up in Sage to pay expense "
     "reimbursements. None has an open balance; the last was paid 2025-05-08. Ramp syncs each reimbursement to NetSuite as "
     "a vendor bill, with the employee as the vendor. It looks for a vendor named exactly \"[First Name] [Last Name] Ramp\" "
     "and creates one if there is none. Vendors loaded under the Sage name would not match, so each employee would end up "
     "in NetSuite twice. Ramp does not sync reimbursements to NetSuite employee records or expense reports (Ramp, 2026-10-28).",
     "Leave the employees out of the load and let Ramp create its own reimbursement vendors; Ramp is already set up with "
     "one vendor per employee (Anthony, 2026-10-28).",
     f"The {B['employees_out']} MinIO employees are not in the file. Neil Martin (V-01355), in Ramp as a vendor, stays: the "
     "Ramp sync brings him over."),
    ("Subsidiary",
     "Every vendor was assumed to do business mainly with MinIO, Inc. Vendors with INR or SGD currency in Sage were marked "
     "\"Subsidiary to confirm\" in case they belong to MinIO Pvt. Ltd. or MinIO Singapore Pte. Ltd.",
     "MinIO, Inc. for every vendor (Ming, meeting of 2026-10-28).",
     "Primary Subsidiary is MinIO, Inc. on every row; the \"Subsidiary to confirm\" marker is removed."),
    ("Duplicate vendors",
     "31 vendor names appeared more than once in Sage (73 records). The meeting also asked for duplicates by email or "
     "address, not only by name.",
     "Merge intelligently: where records are the same vendor by name, email, address or a combination, the record with the "
     "most recent activity wins; where only the name matches, review, but most likely the most recent activity wins "
     "(meeting of 2026-10-28).",
     f"Of the 31 names, {31 - len(kept_dups) - 1} have no record in the load. In {len(kept_dups)}, one record is in the load and it is the one "
     f"with the most recent activity; the Sage IDs of the others, not loaded, are in Merged Sage IDs ({ids(kept_dups)}). "
     "HostIndia has two records in Ramp (next decision). By email and address, Linux Foundation V00208 (Sage only, same "
     f"address and email) merges into The Linux Foundation V-00443, which has the more recent activity. Vendors that share "
     "only a MinIO mailbox (payables@minio.io, accountspayable@min.io and so on, set on Cigna, Kaiser, Rippling, tax "
     "agencies and others) are separate vendors and are not merged."),
    ("Duplicates inside Ramp",
     "Four vendors are in Ramp twice. In each pair one record was created in Ramp and a second, linked to Sage, a few days "
     "later. The Ramp sync would bring both into NetSuite, so this file cannot merge them.",
     "Keep the records linked to Sage; archive the others in Ramp before the sync (Anthony, 2026-10-28). Nisha to action.",
     f"The {len(archived)} records to archive are marked \"Archive in Ramp before the sync\" in Duplicate Action, listed under "
     f"the record kept, and are not counted in the load: " + "; ".join(
         f"{V(r, 'Company Name')} ({V(r, 'Legacy Vendor ID') or 'Ramp only'}), a duplicate of {V(r, 'Duplicate Action').split('duplicate of ')[1]}"
         for r in archived) + "."),
    ("Moss Adams and Baker Tilly",
     "Moss Adams LLP (V-00453, last used 2025-05-19) and Baker Tilly Advisory Group Parent, LP (V-01326, billing now) share "
     "the email mossadamsar@mossadams.com: Baker Tilly took over Moss Adams in 2025.",
     "Keep both as separate vendors (Anthony, 2026-10-28).",
     "Both are in the load, each flagged \"Shares an email with\" the other."),
    ("Tax IDs and bank details",
     "The review copy for the controller and AP showed tax IDs and bank details in full. SOA needs them for the import.",
     "Send SOA the full version, all in one workbook (Anthony, 2026-10-28).",
     "Tax ID and the bank columns are in full on the Vendors-Draft tab. Share this file only with SOA, the controller and AP."),
    ("Vendor IDs",
     "The preliminary file gave the 376 older V##### and Ramp-created RMP_# Sage IDs new V-##### numbers. Most vendors "
     "come into NetSuite through the Ramp sync, which creates them and sets their ID, so an ID in this file could clash "
     "with the one Ramp sets, and changing IDs could affect Ramp's link to the vendor.",
     "For vendors that come through the Ramp sync, Vendor ID and External ID read SYNCED FROM RAMP, so SOA plans how the "
     "sync and this file fit together rather than loading the column as is. Sage only vendors keep their Sage ID "
     "unchanged; nothing is renumbered (Anthony, 2026-10-28).",
     f"Vendor ID and External ID are SYNCED FROM RAMP on the {len(synced_id)} Ramp vendors, and Source System reads "
     f"\"Synced from Ramp\". The {len(from_file)} Sage only vendors keep their Sage ID, with Source System \"Sage Intacct\". "
     "Every row keeps its original IDs: Legacy Vendor ID and Sage Vendor ID hold the Sage ID, Ramp Vendor ID holds Ramp's "
     "ID."),
    ("Individuals without first and last name",
     "NetSuite needs First Name and Last Name for a vendor that is an individual. Sage has neither for some individuals.",
     "Where the vendor name is a plain first and last name, split it; otherwise leave it for review. Vishnu Daya & Co LLP "
     "is a firm, so Individual is unticked (Anthony, 2026-10-28).",
     f"{len(split_names)} split ({ids(split_names)}). V-01459 Vishnu Daya & Co LLP loads as a company. {len(no_names)} "
     "left blank and listed in the open items."),
]

NOTES = [
    ("Vendor count",
     f"Bridge from Sage's active vendors to the vendors in NetSuite. Sage active vendors (Sage's active vendor list "
     f"export): {n(B['export'])}. Less {len(B['onetime'])} one-time-use vendors, not in Sage's vendor detail export "
     f"({', '.join(f'{i} {v}' for i, v in B['onetime'])}): {n(B['active'])}, of which {B['in_ramp']} are also in Ramp, "
     f"linked by the Sage vendor ID held in Ramp, and {n(B['sage_only'])} are not (Sage only). Less "
     f"{n(B['out_of_scope'])} Sage only vendors with no activity since 2024-10-26 ({B['older']} with older activity only, "
     f"{B['sage_only_none']} with none): {n(B['active'] - B['out_of_scope'])}. Less {B['employees_out']} MinIO employees, Sage only, set up for expense reimbursements, which Ramp now creates its own vendors for: {n(B['active'] - B['out_of_scope'] - B['employees_out'])}. Plus {len(B['ramp_only'])} vendors in Ramp "
     f"that are not in Sage (Ramp only), mostly added since 2025-08: {n(len(rows))}, the rows on the Vendors-Draft tab. Less "
     f"{len(merged)} duplicate merged ({ids(merged)}) and {len(archived)} duplicates in Ramp to archive before the sync: "
     f"{n(len(nsv))} vendors in NetSuite. Of these, {len(by_sync)} come in through the Ramp sync ({B['in_ramp'] - sum(1 for r in archived if 'Ramp matched' in QC(r))} "
     f"Ramp matched and {len(B['ramp_only']) - sum(1 for r in archived if 'Ramp only' in QC(r))} Ramp only) and "
     f"{len(from_file)} are loaded from this file (Sage only). The merged and archived records stay on the tab only so the "
     f"decision can be checked, and must be filtered out before the import. Sage's vendor list report also has "
     f"{B['inactive']} inactive vendors, not counted above and not loaded."),
    ("Bank details",
     "Ramp supersedes Sage. Ramp does not push bank details to NetSuite, so they are imported separately for every vendor. "
     f"Where Ramp holds bank details they are used, including the {len(bank_differ)} vendors where Sage holds a different "
     f"account; Sage's are used only where Ramp has none. Bank details to import: {bank['Ramp'] + bank['Sage']} vendors, "
     f"{bank['Ramp']} from Ramp and {bank['Sage']} from Sage. The bank columns show the details to import in full."),
    ("NetSuite load plan",
     "(1) Ramp syncs its vendors to NetSuite: the Ramp only and Ramp matched vendors, once the duplicates in Ramp are "
     "archived. (2) Sage values Ramp does not push are added to the Ramp matched vendors: Sage Vendor ID, category, default "
     "expense account, 1099 flag, tax ID where Ramp has none. Ramp-created vendors carry no Sage ID, so this file is matched "
     "to them by Ramp vendor name and Ramp Vendor ID and re-keyed to the NetSuite internal ID after the sync. (3) Bank "
     f"details are imported for every vendor, as above. (4) The {len(from_file)} Sage only vendors are loaded from this file."),
    ("Vendor ID",
     f"Vendors synced from Ramp: Vendor ID and External ID read SYNCED FROM RAMP; their ID in NetSuite is whatever the Ramp "
     f"sync sets. Sage only vendors: Vendor ID and External ID are the Sage ID, unchanged. Legacy Vendor ID and Sage Vendor "
     f"ID hold the original Sage ID on every row from Sage; Ramp Vendor ID holds Ramp's own ID on every row from Ramp, and "
     f"is the only ID a Ramp only vendor has. The highest Sage V-##### is {data['last_std']}, so NetSuite's vendor "
     f"auto-numbering, if used, should start at {data['next_id']}."),
    ("Category",
     "Sage vendor types mapped to the template's Vendor Category list: AR reseller to AR Reseller, MinIO employee to MinIO "
     "Employee, Tax to Tax agency, the rest unchanged."),
    ("Addresses",
     "Address 1 is the Sage main address (label Main), else Ramp's check mailing address. Address 2 is the Sage pay-to "
     "address where different (label Remit-to, default billing)."),
    ("Terms", "Sage terms, else Ramp net terms, mapped to the NetSuite list: Sage \"15 days\" becomes Net 15."),
    ("Accounts",
     "Default Expense Account from the Sage default expense account, translated through the Sage# column of COA Draft v1. "
     "Default Payables Account 20100 Accounts Payable - Trade for every vendor."),
    ("Tax ID and 1099",
     "Tax ID from Sage, else Ramp's TIN; shown in full. 1099 Eligible is TRUE where Sage has a Form 1099 type or the vendor "
     "type is 1099 Vendor."),
]

OPEN = [
    f"Nisha: archive the {len(archived)} duplicate vendors in Ramp before the Ramp sync to NetSuite, keeping the record "
    "linked to Sage in each pair: " + "; ".join(
        f"archive {V(r, 'Company Name')} ({V(r, 'Legacy Vendor ID') or 'Ramp only'}), keep {V(r, 'Duplicate Action').split('duplicate of ')[1]}"
        for r in archived) + ".",
    f"Individuals with no first and last name in Sage, which NetSuite requires: {ids(no_names)}. AP to supply the names.",
    f"Sage expense accounts with no COA Draft v1 mapping, so Default Expense Account is blank: " + "; ".join(
        f"{a} ({c} vendor{'s' if c > 1 else ''})" for a, c in sorted(collections.Counter(
            f.split("account ")[1].split(" has")[0] for r in no_coa for f in r["flags"] if f.startswith("Sage expense account")).items()))
    + ". Ming to map them in COA Draft v1.",
    f"Vendor IDs (SOA): the {len(synced_id)} vendors marked SYNCED FROM RAMP get their ID from the Ramp sync. Before the "
    "sync, confirm with Ramp what ID it sets (the vendor name or NetSuite's auto-number) and that it links to each vendor "
    f"by NetSuite internal ID. If vendor auto-numbering is on during the sync, start it at {data['next_id']} or later so it "
    "cannot reuse a Sage V-##### number held by a Sage only vendor.",
    "After the Ramp sync: re-key this file's Ramp matched and Ramp only rows to the NetSuite internal IDs the sync creates, "
    "using Ramp Vendor ID, then load the enrichment values and bank details (SOA).",
]

ns.column_dimensions["A"].width = 5
ns.column_dimensions["B"].width = 26
ns.column_dimensions["C"].width = 55
ns.column_dimensions["D"].width = 55
ns.column_dimensions["E"].width = 60
ns["A1"] = "MinIO vendor load to NetSuite: notes"; ns["A1"].font = Font(bold=True, size=14)
ns["A2"] = (f"Built 2026-10-28 from the Sage and Ramp vendor exports of 2026-10-25 (attached to the Asana task, {ASANA}). "
            f"Full version for SOA, the controller and AP: tax IDs and bank details in full. Do not circulate further. The "
            f"Vendors-Draft tab has {len(rows)} rows: {len(nsv)} vendors in NetSuite and {len(gone)} duplicates marked in "
            "Duplicate Action, which are not loaded. Columns marked * are required. Flags (last column) lists everything "
            "changed or worth checking on each row. Cell colors show where each value comes from:")
ns.merge_cells("A2:E2"); ns["A2"].alignment = wrap_top; ns.row_dimensions[2].height = 64
row = 3
for label, key in (("From Sage", None), ("From Ramp", "ramp"), ("Set by rule or decision", "rule"),
                   ("Proposed", "prop"), ("Required, missing", "miss")):
    c = ns.cell(row=row, column=2, value=label)
    if key:
        c.fill = PatternFill("solid", fgColor=FILL[key]); c.font = Font(color=INK[key], italic=(key == "prop"))
    c.border = grid
    row += 1
row += 1

def heading(text):
    global row
    ns.cell(row=row, column=1, value=text).font = Font(bold=True, size=12)
    row += 1

heading("Decisions")
for k, h in enumerate(["#", "Topic", "Scenario or question", "Decision", "What changed in this file"], 1):
    c = ns.cell(row=row, column=k, value=h); c.font = Font(bold=True); c.fill = head_fill; c.border = grid
row += 1
for i, d in enumerate(DECISIONS, 1):
    for k, v in enumerate((i,) + d, 1):
        c = ns.cell(row=row, column=k, value=v); c.alignment = wrap_top; c.border = grid
        if k == 2:
            c.font = Font(bold=True)
    row += 1
row += 1
heading("How the file was built")
for i, (topic, text) in enumerate(NOTES, 1):
    ns.cell(row=row, column=1, value=i).alignment = wrap_top
    ns.cell(row=row, column=2, value=topic).font = Font(bold=True)
    ns.cell(row=row, column=2).alignment = wrap_top
    ns.cell(row=row, column=3, value=text).alignment = wrap_top
    ns.merge_cells(start_row=row, start_column=3, end_row=row, end_column=5)
    ns.row_dimensions[row].height = max(30, 15 * (len(text) // 150 + 1))
    row += 1
row += 1
heading("Open items")
for i, text in enumerate(OPEN, 1):
    ns.cell(row=row, column=1, value=i).alignment = wrap_top
    ns.cell(row=row, column=2, value=text).alignment = wrap_top
    ns.merge_cells(start_row=row, start_column=2, end_row=row, end_column=5)
    ns.row_dimensions[row].height = max(30, 15 * (len(text) // 180 + 1))
    row += 1

wb.save(OUT)
print(f"{len(rows)} rows ({len(nsv)} vendors in NetSuite: {len(by_sync)} by Ramp sync, {len(from_file)} from this file), "
      f"{len(headers)} columns -> {OUT} (tabs: {wb.sheetnames})")
