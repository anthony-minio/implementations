#!/usr/bin/env python3
# Check the vendor workbook (the file sent for review) against the raw Sage and Ramp files, independently of the build.
# Every cell is traced back to the source record it should come from, so values that shifted between rows or columns show up.
#   python3 tools/validate_load.py   (run from the vendors folder)
import csv, collections, datetime as dt, json, re
import openpyxl

SAGE, REPORT, RAMP = "data/sage-vendors-all-detail-20261025.csv", "data/sage-vendor-list-report-20261025.csv", "data/ramp-vendors-sensitive-20261025.csv"
WB, MASKED, COA = "data/MinIO_Vendor_Load_to_NetSuite_20261028.xlsx", "data/vendor-review.html", "data/coa-sage-map.json"

rows = list(csv.reader(open(SAGE, encoding="utf-8-sig")))
H = rows[0]
fi = {}
for i, c in enumerate(H):
    fi.setdefault(c, i)
sage = {r[0].strip(): r for r in rows[1:] if r and r[0].strip()}
g = lambda r, f: r[fi[f]].strip()
tax_cols = [i for i, c in enumerate(H) if c == "Tax ID"]
bal = {x[0].strip(): x[7].strip() for x in csv.reader(open(REPORT, encoding="utf-8-sig")) if re.match(r"(V-?\d+|RMP_\d+)$", x[0].strip())}
ramp_all = list(csv.DictReader(open(RAMP, encoding="utf-8-sig")))
ramp_by_sage = {r["Sage Intacct ID"].strip(): r for r in ramp_all if r["Sage Intacct ID"].strip()}
ramp_by_id = {r["Vendor ID"]: r for r in ramp_all}
coa = json.load(open(COA))

ws = openpyxl.load_workbook(WB, read_only=True)["Vendors-Draft"]
data = list(ws.iter_rows(values_only=True))
hdr = [str(h).replace(" *", "") if h else "" for h in data[1]]
# Template columns repeat names (two address blocks); index them by position within the header.
pos = collections.defaultdict(list)
for i, h in enumerate(hdr):
    pos[h].append(i)
def col(row, name, n=0):
    v = row[pos[name][n]]
    return "" if v is None else str(v)
recs = data[2:]
issues = collections.defaultdict(list)
def bad(check, key, detail):
    issues[check].append(f"{key}: {detail}")

def d(s):
    for f in ("%m/%d/%Y", "%Y-%m-%d"):
        try:
            return dt.datetime.strptime((s or "")[:10], f).date().isoformat()
        except ValueError:
            pass
    return ""
term = lambda t: "" if not t else "Due on Receipt" if t.strip().lower() == "due on receipt" else re.sub(r"^(?:net\s*)?(\d+)(?:\s*days?)?$", r"Net \1", t.strip(), flags=re.I)
CAT = {"AR reseller": "AR Reseller", "MinIO employee": "MinIO Employee", "Tax": "Tax agency"}

seen_sage, seen_ramp_only = collections.Counter(), collections.Counter()
for r in recs:
    sid, rid = col(r, "Sage Vendor ID"), col(r, "Ramp Vendor ID")
    key = sid or f"Ramp {rid}"
    if sid:
        seen_sage[sid] += 1
        S = sage.get(sid)
        if not S:
            bad("Sage ID not in Sage", key, sid); continue
        R = ramp_by_sage.get(sid)
        expect = {
            "Company Name": g(S, "Vendor name"), "Legacy Vendor ID": sid,
            "Vendor Category": CAT.get(g(S, "Vendor type ID"), g(S, "Vendor type ID")) if g(S, "Vendor type ID") != "CC Vendor" else "",
            "Comments": g(S, "Comments"), "Fax": g(S, "Fax number"),
            "Primary Currency": g(S, "Default currency") or "USD",
            "Ramp Vendor ID": R["Vendor ID"] if R else "", "In Ramp": "Yes" if R else "No",
        }
        for k, v in expect.items():
            if col(r, k) != v:
                bad(f"{k} vs source", key, f"{col(r, k)!r} expected {v!r}")
        vid = col(r, "Vendor ID")
        gone = col(r, "Duplicate Action").startswith(("Merge into", "Archive in Ramp"))
        # Vendors in Ramp: SYNCED FROM RAMP, the Ramp sync sets their ID. Sage only: the Sage ID unchanged (2026-10-28).
        want_id, want_src = ("SYNCED FROM RAMP", "Synced from Ramp") if R else (sid, "Sage Intacct")
        if not gone and vid != want_id:
            bad("Vendor ID", key, f"{vid!r} expected {want_id!r}")
        if col(r, "Source System") != want_src:
            bad("Source System", key, f"{col(r, 'Source System')!r} expected {want_src!r}")
        if col(r, "ExternalID") != vid:
            bad("ExternalID vs Vendor ID", key, col(r, "ExternalID"))
        if col(r, "Individual") != ("TRUE" if g(S, "This vendor is an individual person") == "Yes" and sid != "V-01459" else "FALSE"):
            bad("Individual vs Sage", key, col(r, "Individual"))
        phone = g(S, "Phone number") or (R["Contact phone number"].strip().lstrip("'") if R else "")
        if col(r, "Phone") != phone:
            bad("Phone vs source", key, f"{col(r, 'Phone')!r} expected {phone!r}")
        email = g(S, "Email address") or (R["Contact email"].strip() if R else "")
        if col(r, "Email") != email:
            bad("Email vs source", key, f"{col(r, 'Email')!r} expected {email!r}")
        if g(S, "Address line 1"):
            a1 = (g(S, "Address line 1"), g(S, "City"), g(S, "State or territory"), g(S, "ZIP or postal code"), g(S, "Country"))
        elif R and R["Check mailing address line 1"].strip():
            a1 = (R["Check mailing address line 1"].strip(), R["Check mailing city"].strip(" ,"), R["Check mailing state/province/region"].strip(),
                  R["Check mailing zip/postal code"].strip(), None)
        else:
            a1 = ("", "", "", "", g(S, "Country"))
        got = (col(r, "Address 1", 0), col(r, "City", 0), col(r, "Province/State", 0), col(r, "Postal Code/Zip", 0), col(r, "Country", 0))
        if any(e is not None and gv != e for gv, e in zip(got, a1)):
            bad("Address 1 vs source", key, f"{got} expected {a1}")
        legal = (R["Vendor legal name"].strip() if R and R["Vendor legal name"].strip() else g(S, "1099 name"))
        if col(r, "Legal Name") != legal:
            bad("Legal Name vs source", key, f"{col(r, 'Legal Name')!r} expected {legal!r}")
        t = term(g(S, "Term")) or (term(R["Net payment terms (days)"]) if R else "")
        if col(r, "Terms") != t:
            bad("Terms vs source", key, f"{col(r, 'Terms')!r} expected {t!r}")
        exp = g(S, "Default expense account")
        e = coa.get(exp.split()[0].split("-")[0], "") if exp else ""
        if col(r, "Default Expense Account") != e:
            bad("Expense account vs COA map", key, f"{col(r, 'Default Expense Account')!r} expected {e!r}")
        tax = next((S[i].strip() for i in tax_cols if S[i].strip()), "") or (R["Taxpayer identification number (TIN)"].strip() if R else "")
        if col(r, "Tax ID") != tax:
            bad("Tax ID vs source", key, "differs")
        if (col(r, "1099 Eligible") == "TRUE") != bool(g(S, "Form 1099 type") or g(S, "Vendor type ID") == "1099 Vendor"):
            bad("1099 vs Sage", key, col(r, "1099 Eligible"))
        # bank: Ramp first, else Sage
        if R and R["Default ACH account number"].strip():
            eb = ("Ramp", R["Default ACH routing number"].strip(), R["Default ACH account number"].strip())
        elif g(S, "ACH account number"):
            eb = ("Sage", g(S, "ACH bank routing number"), g(S, "ACH account number"))
        else:
            eb = ("", "", "")
        gb = (col(r, "Bank Source"), col(r, "Bank Routing Number"), col(r, "Bank Account Number"))
        if gb != eb:
            bad("Bank details vs source", key, f"source {gb[0]!r} expected {eb[0]!r}" + ("" if gb[1:] == eb[1:] else ", numbers differ"))
        rv = {"Sage type": g(S, "Vendor type ID"), "Sage last bill": d(g(S, "Last bill created on")),
              "Sage last payment": d(g(S, "Last payment made on")),
              "Sage balance": f"{float(bal[sid].replace(',', '')):,.2f}" if bal.get(sid) and float(bal[sid].replace(",", "")) else "",
              "Ramp created": R["Creation date"][:10] if R else "", "Ramp last spend": R["Last spend date"] if R else "",
              "Ramp category": R["Ramp category"] if R else "", "Ramp owner": R["Owner name"] if R else ""}
        for k, v in rv.items():
            if col(r, k) != v:
                bad(f"{k} vs source", key, f"{col(r, k)!r} expected {v!r}")
    else:
        seen_ramp_only[rid] += 1
        R = ramp_by_id.get(rid)
        if not R or R["Sage Intacct ID"].strip():
            bad("Ramp only row not a Ramp-only vendor", key, rid); continue
        if not col(r, "Duplicate Action").startswith("Archive in Ramp") and (col(r, "Vendor ID"), col(r, "ExternalID")) != ("SYNCED FROM RAMP", "SYNCED FROM RAMP"):
            bad("Vendor ID", key, f"{col(r, 'Vendor ID')!r} expected 'SYNCED FROM RAMP'")
        if col(r, "Source System") != "Synced from Ramp":
            bad("Source System", key, col(r, "Source System"))
        name = (R["Display name"] or R["Vendor name"]).strip()
        for k, v in (("Company Name", name), ("Email", R["Contact email"].strip()), ("Tax ID", R["Taxpayer identification number (TIN)"].strip()),
                     ("Ramp created", R["Creation date"][:10]), ("Ramp last spend", R["Last spend date"])):
            if col(r, k) != v:
                bad(f"{k} vs Ramp (Ramp only)", key, "differs" if k == "Tax ID" else f"{col(r, k)!r} expected {v!r}")

# Scope (Ming, 2026-10-28): every Sage vendor in Ramp, and Sage only vendors billed or paid since 2024-10-26.
# Sage only MinIO employees are left out: Ramp creates its own reimbursement vendors (2026-10-28).
active = {i for i, S in sage.items() if i in ramp_by_sage or (max(d(g(S, "Last bill created on")), d(g(S, "Last payment made on"))) >= "2024-10-26"
          and g(S, "Vendor type ID") != "MinIO employee")}
# Records decided out of the load (2026-10-28) are not in the file: a merged duplicate, duplicates in Ramp and Ramp vendors
# no longer used, all archived in Ramp before the sync.
MERGED = {"V00208": "V-00443"}
OUT_SAGE = set(MERGED) | {"RMP_5", "V-01200", "V-01355"}
OUT_RAMP = {"6a461530-6faa-45d1-9cc9-1789847d4bb3", "12040107-b5aa-48d7-8d79-779dc1766034", "c3639aaa-cbca-42f5-aae1-9378d0ceb6b3"}
for r in recs:
    if col(r, "Primary Subsidiary") != "MinIO, Inc.":
        bad("Subsidiary not MinIO, Inc.", col(r, "Sage Vendor ID") or col(r, "Ramp Vendor ID"), col(r, "Primary Subsidiary"))
    if col(r, "Duplicate Action").startswith(("Merge into", "Archive in Ramp")):
        bad("Record not loaded is in the load file", col(r, "Sage Vendor ID") or col(r, "Ramp Vendor ID"), col(r, "Duplicate Action"))
# No two vendors in NetSuite with the same name; Merged Sage IDs are same-name Sage records left out of the load.
nk = lambda t: re.sub(r"[^a-z0-9]", "", t.lower())
live = [r for r in recs if not col(r, "Duplicate Action").startswith(("Merge into", "Archive in Ramp"))]
for k, c in collections.Counter(nk(col(r, "Company Name")) for r in live).items():
    if c > 1:
        bad("Same name twice in NetSuite", k, c)
for r in recs:
    for m in filter(None, col(r, "Merged Sage IDs").split(", ")):
        if m in active and m not in MERGED:
            bad("Merged Sage ID is in the load", col(r, "Sage Vendor ID"), m)
        elif m not in active and (m not in sage or nk(g(sage[m], "Vendor name")) != nk(col(r, "Company Name"))):
            bad("Merged Sage ID not a same-name record", col(r, "Sage Vendor ID"), m)
for m, k in MERGED.items():
    if not any(col(r, "Sage Vendor ID") == k and m in col(r, "Merged Sage IDs").split(", ") for r in recs):
        bad("Merged record not listed on the vendor kept", k, m)
active -= OUT_SAGE
nl = list(openpyxl.load_workbook(WB, read_only=True)["Not Loaded"].iter_rows(values_only=True))[2:]
nl_keys = {str(x[pos["Sage Vendor ID"][0]] or x[pos["Ramp Vendor ID"][0]]) for x in nl}
if nl_keys != OUT_SAGE | OUT_RAMP or len(nl) != len(nl_keys):
    bad("Not Loaded tab", "-", f"extra {sorted(nl_keys - OUT_SAGE - OUT_RAMP)}, missing {sorted(OUT_SAGE | OUT_RAMP - nl_keys)}")
if set(seen_sage) != active:
    bad("Row set", "Sage", f"missing {sorted(active - set(seen_sage))[:5]}, extra {sorted(set(seen_sage) - active)[:5]}")
for i, n in seen_sage.items():
    if n > 1:
        bad("Row set", i, f"appears {n} times")
ro = {r["Vendor ID"] for r in ramp_all if not r["Sage Intacct ID"].strip()} - OUT_RAMP
if set(seen_ramp_only) != ro:
    bad("Row set", "Ramp only", f"missing {len(ro - set(seen_ramp_only))}, extra {len(set(seen_ramp_only) - ro)}")
vids = collections.Counter(col(r, "Vendor ID") for r in recs if col(r, "Vendor ID") and col(r, "Vendor ID") != "SYNCED FROM RAMP")
if any(n > 1 for n in vids.values()):
    bad("Unique Vendor ID", "-", [k for k, n in vids.items() if n > 1][:5])

# The masked page must match the workbook row for row, apart from masked values.
s = open(MASKED, encoding="utf-8").read()
page = json.loads(s[s.index("const DATA = ") + 13:s.index(";\nconst $")])
RVI, SVI = (hdr.index("Ramp Vendor ID"), hdr.index("Sage Vendor ID"))
pkey = lambda cells: cells[SVI] or cells[RVI]
by_page = {pkey([c[0] for c in pr["cells"]]): pr for pr in page["rows"]}
missing = [col(wr, "Sage Vendor ID") or col(wr, "Ramp Vendor ID") for wr in recs if pkey(["" if v is None else str(v) for v in wr]) not in by_page]
if missing:
    bad("Masked page rows", "-", f"workbook rows not on the page: {missing[:5]}")
else:
    for wr in recs:
        pr = by_page[pkey(["" if v is None else str(v) for v in wr])]
        vals = [c[0] for c in pr["cells"]] + pr["review"]
        for j, v in enumerate(vals):
            w = "" if wr[j] is None else str(wr[j])
            if v != w and "•" not in v:
                bad("Masked page vs workbook", col(wr, "Sage Vendor ID") or col(wr, "Ramp Vendor ID"), f"{hdr[j]}: {v!r} vs {w!r}")
                break

print(f"Checked {len(recs)} workbook rows against {len(sage)} Sage vendors, {len(bal)} Sage report rows, {len(ramp_all)} Ramp vendors.")
for k in sorted(issues):
    print(f"\n{k}: {len(issues[k])}")
    for e in issues[k][:6]:
        print("   ", e)
if not issues:
    print("No issues.")
